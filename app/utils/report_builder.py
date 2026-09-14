import os
from openpyxl import Workbook
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from app.config import Config


def ensure_dir(path):
    os.makedirs(os.path.dirname(path), exist_ok=True)


def build_excel(report_no: str, title: str, content: str) -> str:
    path = os.path.join(Config.REPORT_DIR, f"{report_no}.xlsx")
    ensure_dir(path)

    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    ws["A1"] = "报告编号"
    ws["B1"] = report_no
    ws["A2"] = "标题"
    ws["B2"] = title
    ws["A3"] = "内容"
    ws["B3"] = content
    wb.save(path)
    return path


def build_pdf(report_no: str, title: str, content: str) -> str:
    path = os.path.join(Config.REPORT_DIR, f"{report_no}.pdf")
    ensure_dir(path)

    c = canvas.Canvas(path, pagesize=A4)
    c.setFont("Helvetica-Bold", 16)
    c.drawString(80, 780, f"Report: {report_no}")
    c.setFont("Helvetica", 12)
    c.drawString(80, 750, f"Title: {title}")
    c.drawString(80, 720, f"Content: {content[:80]}")
    c.save()
    return path